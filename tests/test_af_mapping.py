import csv
import json
import re
import tempfile
import unittest
from pathlib import Path

import openpyxl

import build_data


class AfMappingTests(unittest.TestCase):
    def test_reads_nbos_name_from_a_sheet_with_title_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "af-map.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "PSEO"
            ws.append(["AF코드 사용 내역 - 브랜드광고"])
            ws.append([])
            ws.append(["AF코드 정보"])
            ws.append(["매체코드", "매체코드명(시트기재)", "매체코드명(NBOS등록)"])
            ws.append(["PSBRD1", "닥스 가디건 추천", "1_brand_닥스 가디건 추천"])
            wb.save(path)

            mapping, aliases = build_data.read_af_map(path)

            self.assertEqual(mapping["PSBRD1"]["n"], "닥스 가디건")
            self.assertEqual(
                mapping["PSBRD1"]["nbos"], "1_brand_닥스 가디건 추천"
            )
            self.assertEqual(aliases["1_brand_닥스 가디건 추천"], "PSBRD1")

    def test_performance_rows_accept_nbos_registered_name_as_af_code(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "performance.csv"
            with path.open("w", encoding="utf-8", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(["", "AF코드", "20260710"])
                writer.writerow(["PV", "1_brand_닥스 가디건 추천", "7"])

            data, _ = build_data.read_perf(
                path, {"1_brand_닥스 가디건 추천": "PSBRD1"}
            )

            self.assertEqual(data[("20260710", "PSBRD1")]["pv"], 7)

    def test_performance_rows_accept_single_metric_uv_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "pSEO_UV.csv"
            path.write_text(
                "\t유입_연월일(YYYYMMDD)\n"
                "AF코드\t20260710\n"
                "총합계\t5\n"
                "PSBRD1\t3\n",
                encoding="utf-16",
            )

            data, totals = build_data.read_perf(path)

            self.assertEqual(data[("20260710", "PSBRD1")]["uv"], 3)
            self.assertEqual(totals["20260710"]["uv"], 5)

    def test_build_uses_nbos_alias_without_adding_unused_sheet_codes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            perf = root / "performance.csv"
            ledger = root / "ledger.xlsx"
            af_map = root / "af-map.xlsx"
            output = root / "data.js"

            with perf.open("w", encoding="utf-8", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(["", "AF코드", "20260710"])
                writer.writerow(["PV", "1_brand_시트 이름 추천", "7"])

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "전체대장"
            ws.append([
                "차수", "키워드", "AF코드", "slug", "현재 운영중",
                "업데이트 날짜", "최신 반영", "미반영 분류", "미반영 사유",
            ])
            ws.append(["1차", "대장 이름 추천", "PSBRD1", "ledger-slug", "Y", "", "Y", "", ""])
            wb.save(ledger)
            wb.close()

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "PSEO"
            ws.append(["AF코드 사용 내역"])
            ws.append([])
            ws.append(["AF코드 정보"])
            ws.append(["매체코드", "매체코드명(시트기재)", "매체코드명(NBOS등록)"])
            ws.append(["PSBRD1", "시트 이름 추천", "1_brand_시트 이름 추천"])
            ws.append(["PSBRD2", "미사용 이름 추천", "2_brand_미사용 이름 추천"])
            wb.save(af_map)
            wb.close()

            build_data.build(perf, ledger, output, af_map)
            payload = output.read_text(encoding="utf-8")
            data = json.loads(payload.removeprefix("window.PSEO_DATA = ").removesuffix(";\n"))

            self.assertEqual([c["c"] for c in data["codes"]], ["PSBRD1"])
            self.assertEqual(data["codes"][0]["n"], "대장 이름")
            self.assertEqual(data["codes"][0]["nbos"], "1_brand_시트 이름 추천")
            self.assertEqual(data["rows"], [["2026-07-10", 0, 7, 0]])

    def test_refresh_embedded_html_adds_latest_nbos_alias_without_renaming_ledger(self):
        self.assertTrue(hasattr(build_data, "refresh_embedded_html"))
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "dashboard.html"
            af_map = root / "af-map.xlsx"
            output = root / "dashboard-refreshed.html"
            data = {
                "codes": [{"c": "PSBRD1", "n": "대장 이름", "src": "ledger"}],
                "rows": [],
                "source": {"map": "ledger.xlsx"},
            }
            source.write_text(
                '<script>window.PSEO_DATA = '
                + json.dumps(data, ensure_ascii=False)
                + ';</script>',
                encoding="utf-8",
            )

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "PSEO"
            ws.append(["AF코드 사용 내역"])
            ws.append([])
            ws.append(["AF코드 정보"])
            ws.append(["매체코드", "매체코드명(시트기재)", "매체코드명(NBOS등록)"])
            ws.append(["PSBRD1", "시트 이름 추천", "1_brand_시트 이름 추천"])
            wb.save(af_map)
            wb.close()

            result = build_data.refresh_embedded_html(
                source, af_map, output, source_name="최신 AF 시트"
            )
            refreshed = output.read_text(encoding="utf-8")
            payload = re.search(
                r"window\.PSEO_DATA\s*=\s*(\{[\s\S]*?\});", refreshed
            )
            self.assertIsNotNone(payload)
            updated = json.loads(payload.group(1))

            self.assertEqual(result, {"matched": 1, "missing": 0, "sheet_codes": 1})
            self.assertEqual(updated["codes"][0]["n"], "대장 이름")
            self.assertEqual(
                updated["codes"][0]["nbos"], "1_brand_시트 이름 추천"
            )
            self.assertEqual(updated["source"]["map"], "ledger.xlsx + 최신 AF 시트")


if __name__ == "__main__":
    unittest.main()
